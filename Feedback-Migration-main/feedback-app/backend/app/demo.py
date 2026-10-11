"""Demo mode: an in-memory stand-in for the TrainingFeedback database.

Enabled with FEEDBACK_DEMO_MODE=true. It lets the application be explored
without access to SQL Server or the StationH login service. Every stored
procedure the application calls has a small Python equivalent here working on
sample data; changes last until the server restarts and no mail is sent.

This is an approximation for demonstrations. The real behaviour is defined by
the stored procedures in the database, not by this module.
"""

from __future__ import annotations

import threading
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any, Callable

from .db import ProcResult, Table

DEMO_PASSWORD = "demo"
FRESHER_COURSE = "FTP_001"
QUESTIONS = {21: "Training Administration", 22: "Training Design And Methodology", 23: "Training Delivery"}


def _fmt(value: date) -> str:
    return value.strftime("%d %b %Y")


def _short(value: date) -> str:
    return f"{value.month}/{value.day}/{value.year}"


@dataclass
class Employee:
    emp_id: str
    name: str
    department: str = "Delivery"
    supervisor: str = ""
    pm: str = ""

    @property
    def mail(self) -> str:
        return self.name.lower().replace(" ", ".") + "@example.com"


@dataclass
class Session:
    session_id: str
    course_id: str
    start: date
    end: date
    location: str
    trainers: list[str]


# login name -> (employee id, description shown on the login page)
DEMO_USERS = {
    "admin": ("10001", "Administrator: every menu and report"),
    "employee": ("20001", "Trainee with feedback to record"),
    "supervisor": ("30001", "Supervisor with trainees to rate"),
    "manager": ("40001", "Project manager with ratings to review"),
}


class DemoDatabase:
    def __init__(self) -> None:
        self.lock = threading.RLock()
        today = date.today()

        people = [
            Employee("10001", "Asha Rao", "Learning & Development"),
            Employee("20001", "Ravi Kumar", supervisor="30001"),
            Employee("20002", "Neha Singh", supervisor="30001"),
            Employee("20003", "Arjun Patel", supervisor="30002"),
            Employee("20004", "Priya Menon", supervisor="30002"),
            Employee("30001", "Meera Nair", pm="40001"),
            Employee("30002", "Rahul Verma", pm="40001"),
            Employee("40001", "Vikram Shah"),
            Employee("50001", "Suresh Iyer", "Learning & Development"),
            Employee("50002", "Kavita Joshi", "Learning & Development"),
        ]
        self.employees = {e.emp_id: e for e in people}
        self.admins = {"10001": 1}  # employee id -> IsSuperAdmin

        self.courses = {
            "JAVA_101": "Core Java",
            "SQL_201": "SQL Server Fundamentals",
            FRESHER_COURSE: "Fresher Training Program",
        }
        self.course_types = ["Technical", "Behavioral", "Process"]
        self.vendors = {"Acme Learning": "101"}
        # (trainer label, type, vendor, course)
        self.trainers = [
            ("Suresh Iyer(50001)", "I", None, "JAVA_101"),
            ("Kavita Joshi(50002)", "I", None, "JAVA_101"),
            ("Kavita Joshi(50002)", "I", None, "SQL_201"),
            ("Kavita Joshi(50002)", "I", None, FRESHER_COURSE),
            ("Daniel Reed(ACM1)", "E", "Acme Learning", "JAVA_101"),
        ]

        java = self._new_session("JAVA_101", "Pune", today - timedelta(days=20), today - timedelta(days=18),
                                 ["Suresh Iyer(50001)", "Kavita Joshi(50002)"])
        sql = self._new_session("SQL_201", "Chennai", today - timedelta(days=10), today - timedelta(days=9),
                                ["Kavita Joshi(50002)"])
        fresher = self._new_session(FRESHER_COURSE, "Chennai", today - timedelta(days=120),
                                    today - timedelta(days=30), ["Kavita Joshi(50002)"])
        self.sessions = {s.session_id: s for s in (java, sql, fresher)}

        self.nominations: list[tuple[str, str]] = (
            [(java.session_id, e) for e in ("20001", "20002", "20003", "20004")]
            + [(sql.session_id, e) for e in ("20001", "20002")]
            + [(fresher.session_id, e) for e in ("20001", "20002", "20003")]
        )
        self.supervisor_list: list[tuple[str, str]] = [
            (fresher.session_id, "30001"),
            (fresher.session_id, "30002"),
        ]

        # trainee feedback: one row per employee and trainer
        self.feedback: list[dict[str, Any]] = []
        self.comments: list[dict[str, Any]] = []
        for emp, trainer, scores, comment, valid in (
            ("20002", "Suresh Iyer(50001)", (5, 4, 5), "Very clear explanations and useful exercises.", 1),
            ("20002", "Kavita Joshi(50002)", (0, 0, 4), "Good pace, would like more hands-on time.", 1),
            ("20003", "Suresh Iyer(50001)", (4, 4, 4), "The lab machines were slow on the first day.", 0),
            ("20003", "Kavita Joshi(50002)", (0, 0, 5), "Excellent session on collections.", 1),
            ("20004", "Suresh Iyer(50001)", (3, 4, 4), "Course material could be shared in advance.", 1),
        ):
            self._add_feedback(java.session_id, "JAVA_101", emp, trainer, *scores, comment, valid)

        # supervisor ratings for the fresher programme
        self.sup_feedback: list[dict[str, Any]] = [
            {
                "session": fresher.session_id, "course": FRESHER_COURSE, "supervisor": "30002",
                "subordinate": "20003", "tech": "4", "soft": "3", "performance": "4", "overall": "4",
                "feedback": "Picked up the project stack quickly; should speak up more in reviews.",
                "work_area": "API development", "status": "1", "remarks": "", "date": today - timedelta(days=3),
            }
        ]

        signature = "<br/><br/>Regards,<br/><signature><br/><replymailid>"
        self.mail_templates = {
            1: {"Subject": "Feedback request for <Course Name>", "CccMail": "", "BccMail": "",
                "Body": "Dear <traineename>,<br/>Please record your feedback for <coursename> held from "
                        "<startdate> to <EndDate> by <Singlerorplural> <trainername>." + signature},
            2: {"Subject": "Training effectiveness rating request", "CccMail": "", "BccMail": "",
                "Body": "Dear <Supervisor Name>,<br/>Please rate your <Singlerorplural> who completed the "
                        "Fresher Training Program." + signature},
            3: {"Subject": "Training effectiveness ratings submitted", "CccMail": "", "BccMail": "",
                "Body": "Dear < PM Name>,<br/><Reporting Manager Name> has submitted ratings for "
                        "<Trainee Details>Please review them." + signature},
            4: {"Subject": "Training effectiveness ratings <status>", "CccMail": "", "BccMail": "",
                "Body": "Dear <Reporting Manager Name>,<br/><PM Name> <status> <Trainee Details>"
                        "Remarks: <remarks><table><row></table>" + signature},
        }
        self.mail_log: list[dict[str, Any]] = []

        # Excel Data Upload: import procedure name -> stored rows
        self.uploads: dict[str, list[dict[str, Any]]] = {}

    # --- helpers --------------------------------------------------------------

    def _new_session(self, course_id: str, location: str, start: date, end: date, trainers: list[str]) -> Session:
        session_id = f"{location[:3]}_{course_id}_{_short(start)}To{_short(end)}"
        return Session(session_id, course_id, start, end, location, list(trainers))

    def _add_feedback(self, session, course, emp, trainer, r1, r2, r3, comment, valid=0) -> None:
        self.feedback.append({"session": session, "course": course, "emp": emp, "trainer": trainer,
                              "r1": r1, "r2": r2, "r3": r3, "common": 1 if r1 else 0})
        self.comments.append({"session": session, "emp": emp, "trainer": trainer,
                              "comment": comment, "valid": valid})

    def person(self, emp_id: str | None) -> Employee:
        emp_id = (emp_id or "").strip()
        return self.employees.get(emp_id) or Employee(emp_id, f"Employee {emp_id}")

    def _nominated(self, session_id: str) -> list[str]:
        return [e for s, e in self.nominations if s == session_id]

    def _subordinates(self, session_id: str, supervisor: str) -> list[Employee]:
        return [self.person(e) for e in self._nominated(session_id) if self.person(e).supervisor == supervisor]

    def _rated_trainers(self, session_id: str, emp: str) -> set[str]:
        return {f["trainer"] for f in self.feedback if f["session"] == session_id and f["emp"] == emp}

    # --- sign-in --------------------------------------------------------------

    def usp_GetFiveDigitEmpID(self, p):
        login = str(p["empid"]).strip().lower()
        emp_id = DEMO_USERS.get(login, (login if login in self.employees else "",))[0]
        return Table(["EMPLID"], [[emp_id]] if emp_id else [])

    def usp_IsAdmin(self, p):
        emp = str(p["empid"])
        return Table(["IsSuperAdmin"], [[self.admins[emp]]] if emp in self.admins else [])

    def usp_getindividual(self, p):
        emp = self.person(p["empid"])
        if p.get("Flag") == 1:
            if emp.emp_id not in self.employees:
                return Table(["AdminName", "ReplyMailID"], [])
            return Table(["AdminName", "ReplyMailID"], [[emp.name, emp.mail]])
        columns = ["TRAINEENAME", "courseid", "TrainerID", "location", "Employee_id", "startdate", "enddate", "TOEMAILID"]
        session = self.sessions.get(p.get("sessionid") or "")
        if not session or emp.emp_id not in self._nominated(session.session_id):
            return Table(columns, [])
        return Table(columns, [[emp.name, session.course_id, t, session.location, emp.emp_id,
                                _fmt(session.start), _fmt(session.end), emp.mail] for t in session.trainers])

    def usp_getindividual_Supervisor(self, p):
        emp = self.person(p["empid"])
        columns = ["Supervisor", "location", "Employee_id", "startdate", "enddate"]
        session = self.sessions.get(p.get("sessionid") or "")
        if not session:
            return Table(columns, [])
        return Table(columns, [[emp.name, session.location, emp.emp_id, _fmt(session.start), _fmt(session.end)]])

    # --- lookups --------------------------------------------------------------

    def usp_getYear(self, p):
        years = sorted({s.start.year for s in self.sessions.values()}, reverse=True)
        return Table(["code_id"], [[y] for y in years])

    def usp_getAllCourse(self, p):
        return Table(["courseid", "course"], sorted(([c, n] for c, n in self.courses.items()), key=lambda r: r[1]))

    def usp_getCourseType(self, p):
        return Table(["CourseTypeID", "CourceTypeName"], [[i + 1, t] for i, t in enumerate(self.course_types)])

    def usp_getSessionCourseWise(self, p):
        rows = [[s.session_id] for s in self.sessions.values()
                if s.course_id == p["courseid"] and s.start.year == int(p["year"])]
        return Table(["sessionid"], rows)

    def usp_getSessionSessionIdWise(self, p):
        columns = ["PreparedDate", "sessionid", "startdate", "enddate", "max_participants", "Trainer Name"]
        session = self.sessions.get(p["sessionid"])
        if not session:
            return Table(columns, [])
        count = len(self._nominated(session.session_id))
        return Table(columns, [[_fmt(date.today()), session.session_id, _fmt(session.start), _fmt(session.end),
                                count, t] for t in session.trainers])

    def usp_getTrainersForSessionid(self, p):
        session = self.sessions.get(p["sessionid"])
        trainers = list(session.trainers) if session else []
        if p.get("empid"):
            done = self._rated_trainers(p["sessionid"], p["empid"])
            trainers = [t for t in trainers if t not in done]
        return Table(["TrainerID"], [[t] for t in trainers])

    def usp_getEmpIdForSessionid(self, p):
        return Table(["Employee_id"], [[e] for e in self._nominated(p["sessionid"])])

    def usp_SelectVendorTrainerDetails(self, p):
        return Table(["VendorName"], [[v] for v in self.vendors])

    def usp_SelectTrainerDetails(self, p):
        rows = [[label] for label, kind, vendor, course in self.trainers
                if kind == p["TrainerType"] and course == p["courseid"]
                and (kind == "I" or vendor == p.get("vendorName"))]
        return Table(["Trainer"], rows)

    def usp_getHexaViewEmployees(self, p):
        emp_id, name = (p.get("empid") or "").strip(), (p.get("EmplName") or "").strip().lower()
        rows = [[e.emp_id, e.name, e.mail] for e in self.employees.values()
                if (emp_id and e.emp_id == emp_id) or (name and name in e.name.lower())
                or (not emp_id and not name)]
        return Table(["EMPLOYEE ID", "EMPLOYEE NAME", "EMAIL"], sorted(rows, key=lambda r: r[1]))

    # --- masters --------------------------------------------------------------

    def usp_InsertCourse(self, p):
        id_exists = int(p["courseId"] in self.courses)
        name_exists = int(p["course"].lower() in (n.lower() for n in self.courses.values()))
        if not id_exists and not name_exists:
            self.courses[p["courseId"]] = p["course"]
        # Like the procedure: a flag is 1 when taken and left NULL otherwise.
        return ProcResult(outputs={"IsCourseIdExist": id_exists or None,
                                   "IsCourseNameExist": name_exists or None,
                                   "RowAffected": int(not id_exists and not name_exists)})

    def usp_InsertVendorDetails(self, p):
        exists = int(p["VendorName"] in self.vendors)
        if not exists:
            self.vendors[p["VendorName"]] = p["PrefixID"]
        return ProcResult(outputs={"IsExistsAlready": exists})

    def usp_InsertTrainerDetails(self, p):
        if p["TrainerType"] == "E":
            prefix = self.vendors.get(p["VendorName"], "EXT")
            number = 1 + sum(1 for _, kind, vendor, _ in self.trainers if kind == "E" and vendor == p["VendorName"])
            trainer_id = f"{prefix}-{number}"
        else:
            trainer_id = p["TrainerId"]
        label = f'{p["TrainerName"]}({trainer_id})'
        entry = (label, p["TrainerType"], p.get("VendorName"), p["course_Id"])
        same_person = [t for t in self.trainers if t[0].split("(")[0] == p["TrainerName"] and t[3] == p["course_Id"]]
        if entry in self.trainers or same_person:
            return ProcResult(outputs={"IsExistsAlready": 1})
        self.trainers.append(entry)
        return ProcResult(outputs={"IsExistsAlready": 0})

    def usp_SelectMailMasterDetails(self, p):
        template = self.mail_templates[int(p["Type"])]
        return Table(["CccMail", "BccMail", "Subject", "Body"],
                     [[template["CccMail"], template["BccMail"], template["Subject"], template["Body"]]])

    def usp_InsertMailDetails(self, p):
        self.mail_templates[int(p["Type"])].update(
            CccMail=p.get("CccMail") or "", Subject=p.get("Subject") or "", Body=p.get("Body") or "")

    # --- sessions -------------------------------------------------------------

    def usp_CheckSessionAvailability(self, p):
        count = sum(1 for s in self.sessions if s.startswith(p["sessionid"]))
        return Table(["SessionCount"], [[count]])

    def usp_InsertSession(self, p):
        if p["sessionid"] in self.sessions:
            return ProcResult(outputs={"IsExistsAlready": 1})
        self.sessions[p["sessionid"]] = Session(p["sessionid"], p["courseid"], p["startdate"], p["enddate"],
                                                p["location"], [])
        return ProcResult(outputs={"IsExistsAlready": 0})

    def usp_InsertSessionTrainerDetails(self, p):
        session = self.sessions[p["sessionid"]]
        if p["TrainerID"] not in session.trainers:
            session.trainers.append(p["TrainerID"])

    # --- nominations and mail -------------------------------------------------

    def usp_getMaxnoofParticipants(self, p):
        has_session = any(s.course_id == p["courseid"] for s in self.sessions.values())
        return Table(["max_participants"], [[500 if has_session else None]])

    def usp_InsertNominationList(self, p):
        target = self.supervisor_list if int(p["Type"]) == 1 else self.nominations
        key = (p["session_id"], str(p["Employee_id"]))
        exists = int(key in target)
        if not exists:
            target.append(key)
        return ProcResult(outputs={"IsExistsAlready": exists, "RowAffected": int(not exists)})

    def usp_insertMasterSupervisorList(self, p):
        return None

    def usp_getSupervisorsForMail(self, p):
        ids = sorted({self.person(e).supervisor for e in self._nominated(p["session_id"])} - {""})
        return Table(["S_EMPLOYEE_ID", "Employee Name"], [[i, self.person(i).name] for i in ids])

    def usp_getEmplForMailsendStatus(self, p):
        rows = [[m["emp"], _fmt(m["date"]), m["status"]] for m in self.mail_log
                if m["session"] == p["session_id"] and m["course"] == p["course_Id"]]
        return Table(["Employee ID", "Nominated Date", "Status"], rows)

    def usp_getMailRelatedDetails(self, p):
        flag = int(p["Flag"])
        emp = self.person(p["Employee_id"])
        session = self.sessions.get(p.get("session_id") or "")

        def template(kind: int) -> Table:
            t = self.mail_templates[kind]
            return Table(["Subject", "Body", "CccMail", "BccMail"],
                         [[t["Subject"], t["Body"], t["CccMail"], t["BccMail"]]])

        if flag == 1:
            return ProcResult([template(1), Table(["AdminName", "ReplyMailID"], [[emp.name, emp.mail]])])
        if flag == 0:
            columns = ["Course Name", "To Mail ID", "Traine Name", "StartDate", "EndDate", "Trainer Name"]
            rows = [[self.courses.get(session.course_id, ""), emp.mail, emp.name, _fmt(session.start),
                     _fmt(session.end), t] for t in (session.trainers if session else [])]
            return ProcResult([template(1), Table(columns, rows)])

        subordinates = self._subordinates(session.session_id, emp.emp_id) if session else []
        pm = self.person(emp.pm)
        saved = [f for f in self.sup_feedback
                 if session and f["session"] == session.session_id and f["supervisor"] == emp.emp_id]
        remarks = saved[0]["remarks"] if saved else ""
        if flag == 2:
            rows = [[emp.mail, emp.name, s.name] for s in subordinates]
            return ProcResult([template(2), Table(["To Mail ID", "Traine Name", "Subordinate Name"], rows)])
        if flag == 3:
            columns = ["SUPERVISOR_MAIL", "To Mail ID", "PM Name", "Supervisor Name", "Subordinate Name", "PM_Remarks"]
            rows = [[emp.mail, pm.mail, pm.name, emp.name, s.name, remarks] for s in subordinates]
            return ProcResult([template(3), Table(columns, rows)])
        columns = ["PM_MAIL", "To Mail ID", "PM Name", "Supervisor Name", "Subordinate Name", "PM_Remarks"]
        rows = [[pm.mail, emp.mail, pm.name, emp.name, s.name, remarks] for s in subordinates]
        return ProcResult([template(4), Table(columns, rows)])

    def usp_sendmail(self, p):
        return None  # nothing is sent in demo mode

    def usp_InsertEmpMailstat(self, p):
        self.mail_log.append({"course": p["course_Id"], "session": p["session_id"], "emp": p["Employee_id"],
                              "date": p["Nominatedate"], "status": p["Status"], "to": p["ToMail"],
                              "subject": p["Subject"]})

    # --- trainee feedback -----------------------------------------------------

    def _training_rows(self, pairs) -> Table:
        columns = ["Course", "Courseid", "sessionid", "startdate", "enddate", "Tot_Feedbacks", "Bal_Feedbacks"]
        rows = [[self.courses.get(s.course_id, s.course_id), s.course_id, s.session_id, _fmt(s.start),
                 _fmt(s.end), total, balance] for s, total, balance in pairs if balance != 0]
        return Table(columns, rows)

    def usp_getTrainingEmpidWise(self, p):
        emp = str(p["empid"])
        pairs = []
        for s in sorted(self.sessions.values(), key=lambda s: s.start):
            if s.course_id != FRESHER_COURSE and emp in self._nominated(s.session_id):
                total = len(s.trainers)
                pairs.append((s, total, total - len(self._rated_trainers(s.session_id, emp))))
        return self._training_rows(pairs)

    def usp_getTrainingSupWise(self, p):
        emp = str(p["empid"])
        pairs = []
        for session_id, supervisor in self.supervisor_list:
            session = self.sessions.get(session_id)
            if supervisor != emp or not session:
                continue
            total = len(self._subordinates(session_id, emp))
            done = sum(1 for f in self.sup_feedback
                       if f["session"] == session_id and f["supervisor"] == emp and f["status"] in ("1", "3"))
            pairs.append((session, total, total - done))
        return self._training_rows(pairs)

    def usp_getReviews(self, p):
        pm = str(p["PM_ID"])
        seen, rows = set(), []
        for f in self.sup_feedback:
            supervisor = self.person(f["supervisor"])
            key = (f["session"], supervisor.emp_id)
            if f["status"] != "1" or supervisor.pm != pm or key in seen:
                continue
            seen.add(key)
            session = self.sessions[f["session"]]
            rows.append([f"{supervisor.name}-{supervisor.emp_id}", self.courses.get(session.course_id, ""),
                         _fmt(session.start), _fmt(session.end), session.session_id, session.course_id])
        return Table(["Supervisors", "Course", "startdate", "enddate", "Session_id", "Course_id"], rows)

    def usp_getFeedBackCount(self, p):
        count = sum(1 for c in self.comments if c["session"] == p["sessionid"] and c["emp"] == p["empid"])
        return Table(["FeedbackCount"], [[count]])

    def usp_Insertfeedbacks_V1(self, p):
        self._add_feedback(p["sessionid"], p["courseid"], str(p["employee"]), p["TrainerId"],
                           int(p["response1"]), int(p["response2"]), int(p["response3"]), p["comments"])

    # --- supervisor ratings ---------------------------------------------------

    def usp_SupervisorFeedback_Details(self, p):
        supervisor, session_id = str(p["empid"]), p["sessionid"]
        session = self.sessions.get(session_id)
        subordinates = self._subordinates(session_id, supervisor) if session else []
        columns = ["Subordinate_Id", "Subordinate_Name", "WorkArea", "TechSkills", "SoftSkills", "Performance",
                   "OvrAllRating", "Overall_Feedback", "Status", "Course", "startdate", "enddate", "Location"]
        saved = []
        for f in self.sup_feedback:
            if f["session"] != session_id or f["supervisor"] != supervisor:
                continue
            if p.get("pm_id") and self.person(supervisor).pm != str(p["pm_id"]):
                continue
            saved.append([f["subordinate"], self.person(f["subordinate"]).name, f["work_area"], f["tech"],
                          f["soft"], f["performance"], f["overall"], f["feedback"], f["status"],
                          self.courses.get(session.course_id, ""), _fmt(session.start), _fmt(session.end),
                          session.location])
        return ProcResult([Table(["Subordinate_Id", "Subordinate_Name"], [[s.emp_id, s.name] for s in subordinates]),
                           Table(columns, saved)])

    def usp_InsertSupFeedbacks(self, p):
        row = next((f for f in self.sup_feedback if f["session"] == p["sessionid"]
                    and f["supervisor"] == p["SupervisorId"] and f["subordinate"] == p["SubOrdId"]), None)
        if row is None:
            row = {"session": p["sessionid"], "course": p["courseid"], "supervisor": p["SupervisorId"],
                   "subordinate": p["SubOrdId"], "remarks": ""}
            self.sup_feedback.append(row)
        row.update(tech=p["TechSkills"], soft=p["SoftSkills"], performance=p["Performance"],
                   overall=p["OvrAllRating"], feedback=p["OverallFeed"], work_area=p["WorkArea"],
                   status=p["Status"], date=date.today())

    def usp_UpdatePMRemarks(self, p):
        for f in self.sup_feedback:
            if f["session"] == p["sessionid"] and f["supervisor"] == p["SupervisorId"]:
                f["status"], f["remarks"] = str(p["Status"]), p["Remarks"]

    # --- reports --------------------------------------------------------------

    def usp_getFeedbackCommentsForSessionid(self, p):
        rows = [[c["emp"], c["comment"], str(c["valid"]), c["trainer"]] for c in self.comments
                if c["session"] == p["sessionid"] and c["comment"]]
        return Table(["Employee_id", "Comments", "CommentsValid", "TrainerId"], sorted(rows, key=lambda r: r[0]))

    def usp_UpdateCommentsvaildDetails(self, p):
        for c in self.comments:
            if c["session"] == p["session_id"] and c["emp"] == p["Employee_id"]:
                c["valid"] = int(p["Commentsvalid"])

    def usp_getFeedbackBalanceForSessionid(self, p):
        columns = ["sessionid", "EMPLOYEES", "Trainers", "Tot_Feedbacks", "Comp_Feedback"]
        session = self.sessions.get(p["sessionid"])
        done = {c["emp"] for c in self.comments if c["session"] == p["sessionid"]}
        if not session or not done:
            return Table(columns, [])
        employees = len(set(self._nominated(session.session_id)))
        return Table(columns, [[session.session_id, employees, len(session.trainers),
                                employees * len(session.trainers), len(done)]])

    def usp_getFeedbackBalanceEmpIdForSessionid(self, p):
        columns = ["Employee ID", "Employee Name", "Department", "Session ID", "Total FeedBack", "Pending FeedBack"]
        session = self.sessions.get(p["sessionid"])
        rows = []
        for emp_id in self._nominated(p["sessionid"]) if session else []:
            emp = self.person(emp_id)
            total = len(session.trainers)
            rows.append([emp.emp_id, emp.name, emp.department, session.session_id, total,
                         total - len(self._rated_trainers(session.session_id, emp_id))])
        return Table(columns, rows)

    def _answers(self, session_id: str, trainer: str) -> dict[int, list[int]]:
        rows = [f for f in self.feedback if f["session"] == session_id]
        delivery = rows if trainer == "OverAll" else [f for f in rows if f["trainer"] == trainer]
        return {21: [f["r1"] for f in rows if f["common"]],
                22: [f["r2"] for f in rows if f["common"]],
                23: [f["r3"] for f in delivery]}

    def usp_getFeedbackForOverallrpt_V1(self, p):
        answers = self._answers(p["sessionid"], p["TrainerId"])
        columns = ["S_No", "QUESTION", "A", "B", "C", "D", "PATTERN"]
        if not answers[23]:
            return Table(columns, [])
        rows = []
        for i, (number, values) in enumerate(answers.items(), start=1):
            counts = Counter(values)
            pattern = "ABCDE"[counts.most_common(1)[0][0] - 1] if counts else ""
            rows.append([i, QUESTIONS[number]] + [counts.get(v, 0) for v in (1, 2, 3, 4)] + [pattern])
        return Table(columns, rows)

    def usp_rATINGForOverallrpt_V1(self, p):
        answers = self._answers(p["sessionid"], p["TrainerId"])
        if not all(answers.values()):
            raise RuntimeError("Divide by zero error encountered.")
        rows = [[i, QUESTIONS[n], 5.0, f"{sum(v) / len(v):.2f}"] for i, (n, v) in enumerate(answers.items(), start=1)]
        return Table(["SNo", "SECTION", "MAXS", "SCORE"], rows)

    def usp_getFeedbackdetQuesansEmployeeid_V1(self, p):
        session_id, trainer, wanted = p["sessionid"], p["TrainerID"], str(p["empid"])
        columns = ["Question No", "QUESTION", "Response", "Session", "Employee Id", "Employee Name"]
        rows = []
        employees = sorted({f["emp"] for f in self.feedback if f["session"] == session_id})
        for emp_id in employees:
            if wanted not in ("All", emp_id):
                continue
            mine = [f for f in self.feedback if f["session"] == session_id and f["emp"] == emp_id]
            common = next((f for f in mine if f["common"]), None)
            delivery = next((f for f in mine if f["trainer"] == trainer), None)
            if delivery is None:
                continue
            name = self.person(emp_id).name
            for number, value in ((21, common and common["r1"]), (22, common and common["r2"]), (23, delivery["r3"])):
                if value:
                    rows.append([str(number), QUESTIONS[number], value, session_id, emp_id, name])
        return Table(columns, rows)

    def usp_getTrnEffectivenessReport(self, p):
        columns = ["serial number", "course", "Session_id", "S_SUPERVISOR_ID", "Supervisor Name", "Subordinate_ID",
                   "Subordinate Name", "Work Area", "Technical Skills", "Soft Skills", "Performance",
                   "Overall Feedback Rating", "Overall Comments", "Status", "Feedback Date", "PM_ID", "PM Name",
                   "PM Remarks"]
        status = {"0": "Saved", "1": "Submitted", "3": "Approved"}
        rows = []
        for f in self.sup_feedback:
            if f["session"] != p["sessionid"]:
                continue
            supervisor = self.person(f["supervisor"])
            pm = self.person(supervisor.pm)
            rows.append([len(rows) + 1, self.courses.get(f["course"], ""), f["session"], supervisor.emp_id,
                         supervisor.name, f["subordinate"], self.person(f["subordinate"]).name, f["work_area"],
                         f["tech"], f["soft"], f["performance"], f["overall"], f["feedback"],
                         status.get(f["status"], f["status"]), _fmt(f["date"]), pm.emp_id, pm.name, f["remarks"]])
        return Table(columns, rows)

    # --- Excel Data Upload ----------------------------------------------------

    def GetSkillportAssessmentReport(self, p):
        columns = ["EmployeeID", "EmployeeName", "PmsCompetencyBasedPlan", "LearningProgramCompletionStatus",
                   "LearningProgramcompletionDate", "AssessmentCompletionStatus", "AssessmentCompletionDate",
                   "Technology", "Role", "Month", "Year"]
        today = date.today()
        rows = [
            ["20001", "Ravi Kumar", "Java Developer L1", "Completed", _fmt(today - timedelta(days=40)),
             "Completed", _fmt(today - timedelta(days=35)), "Java", "Developer", today.month, today.year],
            ["20002", "Neha Singh", "Java Developer L1", "In Progress", "", "Not Started", "", "Java",
             "Developer", today.month, today.year],
        ]
        return Table(columns, rows)

    def _upload_group(self, proc: str) -> tuple[str, str] | None:
        from .routers.skillport import UPLOAD_TYPES

        for check, insert, delete in UPLOAD_TYPES.values():
            for kind, name in (("check", check), ("insert", insert), ("delete", delete)):
                if name == proc:
                    return kind, insert
        return None

    def _upload_proc(self, kind: str, key: str, p: dict) -> Any:
        stored = self.uploads.setdefault(key, [])
        period = (str(p["MONTH"]), str(p["YEAR"]))
        matching = [r for r in stored if (r["MONTH"], r["YEAR"]) == period]
        if kind == "check":
            columns = list(matching[0]) if matching else ["MONTH", "YEAR"]
            return Table(columns, [[r.get(c) for c in columns] for r in matching])
        if kind == "delete":
            self.uploads[key] = [r for r in stored if (r["MONTH"], r["YEAR"]) != period]
            return len(matching)

        # insert: the real procedure reads the workbook inside SQL Server
        from openpyxl import load_workbook

        try:
            # Read from bytes: uploads are stored with an .xls name whatever their format.
            with open(p["FILEPATH"], "rb") as handle:
                workbook = load_workbook(BytesIO(handle.read()), read_only=True, data_only=True)
        except Exception:
            return Table(["Saved"], [[0]])
        rows = [list(r) for r in workbook.worksheets[0].iter_rows(values_only=True) if any(v is not None for v in r)]
        workbook.close()
        if len(rows) < 2:
            return Table(["Saved"], [[0]])
        header = [str(h) if h is not None else f"Column{i + 1}" for i, h in enumerate(rows[0])]
        for values in rows[1:]:
            record = {h: (str(v) if isinstance(v, (datetime, date)) else v) for h, v in zip(header, values)}
            stored.append({**record, "MONTH": period[0], "YEAR": period[1]})
        return Table(["Saved"], [[len(rows) - 1]])

    # --- dispatch -------------------------------------------------------------

    def call(self, name: str, params: dict[str, Any]) -> Any:
        with self.lock:
            group = self._upload_group(name)
            if group:
                return self._upload_proc(group[0], group[1], params)
            handler: Callable[[dict], Any] | None = getattr(self, name, None)
            if handler is None or name.startswith("_"):
                raise NotImplementedError(f"Demo mode has no stand-in for {name}")
            return handler(params)


_database: DemoDatabase | None = None


def database() -> DemoDatabase:
    global _database
    if _database is None:
        _database = DemoDatabase()
    return _database


def reset() -> None:
    global _database
    _database = None


def call_proc(name: str, params: dict[str, Any] | None) -> ProcResult:
    result = database().call(name, dict(params or {}))
    if isinstance(result, ProcResult):
        return result
    if isinstance(result, Table):
        return ProcResult([result])
    return ProcResult()


def call_proc_rowcount(name: str, params: dict[str, Any] | None) -> int:
    result = database().call(name, dict(params or {}))
    return result if isinstance(result, int) else 0


def check_login(userid: str, password: str) -> bool:
    return userid.strip().lower() in DEMO_USERS and password == DEMO_PASSWORD


def login_hints() -> list[dict[str, str]]:
    return [{"username": name, "password": DEMO_PASSWORD, "description": description}
            for name, (_, description) in DEMO_USERS.items()]
