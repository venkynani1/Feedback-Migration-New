"""Excel and PDF downloads."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from html import escape
from io import BytesIO
from typing import Any

from fastapi import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .config import get_settings

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_HEADER_FILL = PatternFill("solid", fgColor="507CD1")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _cell(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if value is None or isinstance(value, (int, float, str, datetime, date)):
        return value
    return str(value)


def workbook_bytes(
    columns: list[str],
    rows: list[list[Any]],
    *,
    heading: list[tuple[str, str]] | None = None,
) -> bytes:
    """A one-sheet workbook; heading is optional label/value lines above the table."""
    workbook = Workbook()
    sheet = workbook.active
    for label, value in heading or []:
        sheet.append([label, value])
        sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
    if heading:
        sheet.append([])

    sheet.append(columns)
    for cell in sheet[sheet.max_row]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        sheet.append([_cell(v) for v in row])

    for i, column in enumerate(columns, start=1):
        longest = max([len(str(column))] + [len(str(r[i - 1] or "")) for r in rows[:200]])
        sheet.column_dimensions[sheet.cell(row=1, column=i).column_letter].width = min(
            max(longest + 2, 10), 60
        )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def download(content: bytes, filename: str, media_type: str) -> Response:
    return Response(
        content,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def feedback_analysis_pdf(
    *,
    course_title: str,
    course_dates: str,
    trainer_name: str,
    participants: str,
    administration: str,
    design: str,
    delivery: str,
    overall: str,
    comments: list[str],
    prepared_by: str,
    prepared_date: str,
) -> bytes:
    """The 'Training Feedback Analysis' sheet (form Hex 2557)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Image,
        ListFlowable,
        ListItem,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    base = getSampleStyleSheet()["Normal"]
    normal = ParagraphStyle("n", parent=base, fontSize=8, leading=11)
    bold = ParagraphStyle("b", parent=normal, fontName="Helvetica-Bold")
    value = ParagraphStyle("v", parent=normal, textColor=colors.HexColor("#006BD7"))
    title = ParagraphStyle("t", parent=bold, alignment=1, fontSize=10)
    right = ParagraphStyle("r", parent=bold, alignment=2)
    blue = colors.HexColor("#006BD7")
    grid = TableStyle(
        [
            ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]
    )

    def p(text: str, style=normal) -> Paragraph:
        return Paragraph(escape(text), style)

    story: list[Any] = []
    logo = get_settings().assets_dir / "HexawareLogo.jpg"
    if logo.exists():
        image = Image(str(logo), width=35 * mm, height=17 * mm, kind="proportional")
        image.hAlign = "LEFT"
        story.append(image)
    story += [p("Hex 2557", right), p("Training Feedback Analysis", title), Spacer(1, 5 * mm)]

    details = Table(
        [
            [p("Course Details", title), "", "", ""],
            [p("Course Title:", bold), p(course_title, value), p("Course Date(s):", bold), p(course_dates, value)],
            [p("Trainer's Name:", bold), p(trainer_name, value), p("No. of Participants:", bold), p(participants, value)],
        ],
        colWidths=["22%", "28%", "22%", "28%"],
    )
    details.setStyle(TableStyle(grid.getCommands() + [("SPAN", (0, 0), (-1, 0))]))
    story += [details, Spacer(1, 5 * mm)]

    scores = Table(
        [
            ["", p("Maximum Score", bold), p("Score", bold)],
            [p("Training Administration:", bold), p("5.0", bold), p(administration, value)],
            [p("Training Design & Methodology:", bold), p("5.0", bold), p(design, value)],
            [p("Trainer Delivery:", bold), p("5.0", bold), p(delivery, value)],
        ],
        colWidths=["40%", "30%", "30%"],
    )
    scores.setStyle(grid)
    story += [scores, Spacer(1, 5 * mm)]

    rating = Table([[p("Overall Rating:", bold), p(overall, value)]], colWidths=["50%", "50%"])
    rating.setStyle(grid)
    story += [rating, Spacer(1, 5 * mm)]

    story.append(
        p(
            "Based on overall feedback received from participants, "
            "record areas of improvements if any:",
            bold,
        )
    )
    if comments:
        story.append(
            ListFlowable(
                [ListItem(p(c), leftIndent=12) for c in comments],
                bulletType="bullet",
                bulletColor=blue,
                leftIndent=12,
            )
        )
    story.append(Spacer(1, 5 * mm))

    prepared = Table(
        [[p("Prepared by:", bold), p(prepared_by, value), p("Date:", bold), p(prepared_date, value)]],
        colWidths=["22%", "28%", "22%", "28%"],
    )
    prepared.setStyle(grid)
    story.append(prepared)

    buffer = BytesIO()
    SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title="Training Feedback Analysis",
    ).build(story)
    return buffer.getvalue()
