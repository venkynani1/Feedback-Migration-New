from __future__ import annotations

from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import db
from app.config import get_settings
from app.main import create_app
from app.routers import reports


# --- sign-in and access control ---------------------------------------------


def test_login_sets_role_and_uses_five_digit_id(client, fake_db):
    fake_db.on("usp_IsAdmin", db.Table(["IsSuperAdmin"], [[1]]))
    response = client.post("/api/auth/login", json={"username": " 1000012345 ", "password": "pw"})
    assert response.json() == {"empId": "12345", "name": "Asha Rao", "role": "super"}
    assert fake_db.calls_to("usp_GetFiveDigitEmpID") == [{"empid": "1000012345"}]
    assert fake_db.calls_to("usp_IsAdmin") == [{"empid": "12345"}]
    assert client.get("/api/auth/me").json()["role"] == "super"

    client.post("/api/auth/logout")
    assert client.get("/api/auth/me").status_code == 401


def test_login_without_five_digit_id_is_refused(client, fake_db):
    fake_db.on("usp_GetFiveDigitEmpID", db.Table(["EMPLID"], []))
    response = client.post("/api/auth/login", json={"username": "x", "password": "pw"})
    assert response.status_code == 403
    assert response.json()["detail"] == "Five digit id not available for this employee"


def test_admin_pages_need_an_admin(user_client):
    assert user_client.get("/api/lookups/courses").status_code == 403
    assert user_client.get("/api/feedback/my").status_code == 200


def test_everything_needs_a_login(client):
    assert client.get("/api/feedback/my").status_code == 401
    assert client.get("/api/reports/defaulters", params={"session_id": "S"}).status_code == 401


def test_posts_without_ajax_header_are_blocked(fake_db):
    bare = TestClient(create_app())
    assert bare.post("/api/auth/login", json={"username": "a", "password": "b"}).status_code == 403


# --- sessions ---------------------------------------------------------------

SESSION = {
    "courseId": "C1",
    "courseName": "Java",
    "location": "Pune",
    "startDate": "2024-03-04",
    "endDate": "2024-03-05",
    "trainers": ["Anil(111)", "Bala(222)"],
}


def test_create_session_inserts_session_then_trainers(admin_client, fake_db):
    fake_db.on("usp_InsertSession", outputs={"IsExistsAlready": 0})
    response = admin_client.post("/api/sessions", json=SESSION)
    assert response.json()["sessionId"] == "Pun_C1_3/4/2024To3/5/2024"
    inserted = fake_db.calls_to("usp_InsertSession")[0]
    assert inserted["startdate"] == date(2024, 3, 4)
    assert inserted["max_participants"] == 500
    assert inserted["createdby"] == "12345"
    assert [c["TrainerID"] for c in fake_db.calls_to("usp_InsertSessionTrainerDetails")] == SESSION["trainers"]


def test_create_another_session_gets_a_numbered_id(admin_client, fake_db):
    fake_db.on("usp_CheckSessionAvailability", db.Table(["SessionCount"], [[2]]))
    fake_db.on("usp_InsertSession", outputs={"IsExistsAlready": 0})
    check = admin_client.post("/api/sessions/check", json=SESSION).json()
    assert check == {"sessionId": "Pun_C1_3/4/2024To3/5/2024", "existingCount": 2}
    created = admin_client.post("/api/sessions", json={**SESSION, "createAnother": True}).json()
    assert created["sessionId"] == "Pun_C1_3/4/2024To3/5/2024_2"


def test_session_rules(admin_client, fake_db):
    fake_db.on("usp_InsertSession", outputs={"IsExistsAlready": 1})
    assert admin_client.post("/api/sessions", json=SESSION).json()["detail"] == "Same session already Exist."
    tomorrow = (date.today() + timedelta(days=1)).isoformat()
    future = admin_client.post("/api/sessions", json={**SESSION, "startDate": tomorrow, "endDate": tomorrow})
    assert future.json()["detail"] == "From Date should be current date or less than current date."
    none = admin_client.post("/api/sessions", json={**SESSION, "trainers": []})
    assert none.json()["detail"] == "Minimum one trainers should select for session."
    assert len(fake_db.calls_to("usp_InsertSessionTrainerDetails")) == 0


# --- masters ----------------------------------------------------------------


def test_external_trainer_needs_vendor_and_gets_no_id(admin_client, fake_db):
    fake_db.on("usp_InsertTrainerDetails", outputs={"IsExistsAlready": 0})
    body = {"courseId": "C1", "trainerType": "E", "trainerId": "Auto generate", "trainerName": "Zed"}
    assert admin_client.post("/api/trainers", json=body).json()["detail"] == "Please select Vendor Name."
    assert admin_client.post("/api/trainers", json={**body, "vendorName": "Acme"}).status_code == 200
    call = fake_db.calls_to("usp_InsertTrainerDetails")[0]
    assert (call["TrainerId"], call["VendorName"]) == (None, "Acme")


def test_mail_master_maps_list_type(admin_client, fake_db):
    fake_db.on("usp_SelectMailMasterDetails", db.Table(["CccMail", "Subject", "Body"], [["c@x.com", "S", "B"]]))
    assert admin_client.get("/api/mail-master", params={"type": 1}).json() == {
        "from": "asha@example.com", "cc": "c@x.com", "subject": "S", "body": "B"
    }
    assert fake_db.calls_to("usp_SelectMailMasterDetails") == [{"Type": 2}]
    bad = admin_client.put("/api/mail-master", json={"type": 0, "cc": "nope", "subject": "s", "body": "b"})
    assert bad.json()["detail"] == "Invalid Email Format"


# --- nominations ------------------------------------------------------------


def _mail_details(params):
    template = db.Table(["Subject", "Body", "CccMail", "BccMail"], [["Sub", "Body", "", ""]])
    if params["Flag"] == 1:
        return db.ProcResult([template, db.Table(["AdminName", "ReplyMailID"], [["Asha", "asha@x.com"]])])
    values = db.Table(
        ["Course Name", "To Mail ID", "Traine Name", "StartDate", "EndDate", "Trainer Name"],
        [["Java", "t@x.com", "Ravi", "a", "b", "Anil(1)"]],
    )
    return db.ProcResult([template, values])


def test_sample_file_parses(admin_client):
    sample = (get_settings().assets_dir / "SampleNominationList.xlsx").read_bytes()
    parsed = admin_client.post("/api/nominations/parse", files={"file": ("list.xlsx", sample)}).json()
    assert parsed["columns"] == ["EMPID", "EMPNAME"]
    assert parsed["rows"][0][0] == "1000035505"
    wrong = admin_client.post("/api/nominations/parse", files={"file": ("list.xls", b"x")})
    assert wrong.json()["detail"].endswith(".xlsx extension.")


def test_nominations_mail_new_employees_only(admin_client, fake_db):
    fake_db.on("usp_getMailRelatedDetails", handler=_mail_details)
    fake_db.on("usp_getMaxnoofParticipants", db.Table(["max_participants"], [[500]]))
    fake_db.on(
        "usp_InsertNominationList",
        handler=lambda p: db.ProcResult(outputs={"IsExistsAlready": int(p["Employee_id"] == "222"), "RowAffected": 1}),
    )
    body = {"listType": 0, "courseId": "C1", "sessionId": "S1", "employeeIds": ["111", " ", "222"]}
    result = admin_client.post("/api/nominations", json=body).json()
    assert result["alreadyNominated"] == ["222"]
    assert [c["to"] for c in fake_db.calls_to("usp_sendmail")] == ["t@x.com"]
    assert [c["Employee_id"] for c in fake_db.calls_to("usp_InsertNominationList")] == ["111", "222"]
    assert fake_db.calls_to("usp_insertMasterSupervisorList") == []


def test_fresher_participants_are_not_mailed_but_supervisors_are(admin_client, fake_db):
    fake_db.on("usp_getMailRelatedDetails", handler=_mail_details)
    fake_db.on("usp_getMaxnoofParticipants", db.Table(["max_participants"], [[500]]))
    fake_db.on("usp_InsertNominationList", outputs={"IsExistsAlready": 0, "RowAffected": 1})
    body = {"listType": 0, "courseId": "FTP_001", "sessionId": "S1", "employeeIds": ["111"]}
    admin_client.post("/api/nominations", json=body)
    assert fake_db.calls_to("usp_sendmail") == []

    admin_client.post("/api/nominations", json={**body, "listType": 1})
    assert len(fake_db.calls_to("usp_sendmail")) == 1
    assert fake_db.calls_to("usp_insertMasterSupervisorList") == [{"session_id": "S1"}]


def test_nomination_limit(admin_client, fake_db):
    fake_db.on("usp_getMailRelatedDetails", handler=_mail_details)
    fake_db.on("usp_getMaxnoofParticipants", db.Table(["max_participants"], [[1]]))
    body = {"listType": 0, "courseId": "C1", "sessionId": "S1", "employeeIds": ["1", "2", "3"]}
    assert admin_client.post("/api/nominations", json=body).json()["detail"].endswith("greater than :1")


# --- trainee feedback -------------------------------------------------------

FEEDBACK = {
    "sessionId": "S1", "courseId": "C1", "trainerId": "Anil(111)",
    "administration": 5, "design": 4, "delivery": 3, "comments": " good ",
}


def test_first_feedback_stores_session_wide_answers(user_client, fake_db):
    fake_db.on("usp_getTrainersForSessionid", db.Table(["TrainerID"], [["Anil(111)"], ["Bala(222)"]]))
    fake_db.on("usp_getFeedBackCount", db.Table(["c"], [[0]]))
    result = user_client.post("/api/feedback", json=FEEDBACK).json()
    assert result == {"remainingTrainers": ["Bala(222)"]}
    saved = fake_db.calls_to("usp_Insertfeedbacks_V1")[0]
    assert (saved["response1"], saved["response2"], saved["response3"], saved["CommonFlag"]) == (5, 4, 3, 1)
    assert (saved["employee"], saved["comments"]) == ("12345", "good")


def test_later_feedback_only_rates_the_trainer(user_client, fake_db):
    fake_db.on("usp_getTrainersForSessionid", db.Table(["TrainerID"], [["Bala(222)"]]))
    fake_db.on("usp_getFeedBackCount", db.Table(["c"], [[1]]))
    body = {**FEEDBACK, "trainerId": "Bala(222)", "administration": None, "design": None}
    assert user_client.post("/api/feedback", json=body).status_code == 200
    saved = fake_db.calls_to("usp_Insertfeedbacks_V1")[0]
    assert (saved["response1"], saved["response2"], saved["response3"], saved["CommonFlag"]) == (0, 0, 3, 0)


def test_feedback_validation(user_client, fake_db):
    fake_db.on("usp_getTrainersForSessionid", db.Table(["TrainerID"], [["Anil(111)"]]))
    fake_db.on("usp_getFeedBackCount", db.Table(["c"], [[0]]))
    post = lambda **changes: user_client.post("/api/feedback", json={**FEEDBACK, **changes}).json()["detail"]
    assert post(trainerId="Someone else") == "Please Select Trainers"
    assert post(comments="  ") == "Please enter your comments"
    assert post(design=None) == "Please answer every question."
    assert fake_db.calls_to("usp_Insertfeedbacks_V1") == []


# --- supervisor ratings and review ------------------------------------------

SUBORDINATES = db.Table(["Subordinate_id", "Name"], [["777", "Trainee One"]])
SAVED = db.Table(
    ["Subordinate_id", "TechSkills", "SoftSkills", "Performance", "OvrAllRating", "Overall_Feedback",
     "WorkArea", "Status", "Course", "startdate", "enddate", "Location", "Subordinate_Name"],
    [["777", "4", "0", "NA", "5", "fine", "QA", "1", "Java", "a", "b", "Pune", "Trainee One"]],
)
ROW = {
    "subordinateId": "777", "workArea": "QA", "techSkills": "4", "softSkills": "3",
    "performance": "NA", "overallRating": "5", "overallFeedback": "fine",
}


def test_supervisor_form_merges_saved_ratings(user_client, fake_db):
    fake_db.on("usp_getindividual_Supervisor", db.Table(
        ["Employee_id", "Supervisor", "startdate", "enddate", "location"], [["12345", "Asha", "a", "b", "Pune"]]))
    fake_db.on("usp_SupervisorFeedback_Details", SUBORDINATES, SAVED)
    form = user_client.get("/api/feedback/supervisor", params={"session_id": "S1", "course_id": "C1"}).json()
    assert form["header"]["dates"] == "a To b"
    assert form["rows"] == [{
        "subordinateId": "777", "name": "Trainee One", "workArea": "QA", "techSkills": "4",
        "softSkills": "", "performance": "NA", "overallRating": "5", "overallFeedback": "fine",
    }]


def test_supervisor_save_and_submit(user_client, fake_db):
    fake_db.on("usp_SupervisorFeedback_Details", SUBORDINATES, SAVED)
    body = {"sessionId": "S1", "courseId": "C1", "submit": False, "rows": [{**ROW, "softSkills": ""}]}
    assert user_client.post("/api/feedback/supervisor", json=body).json()["message"] == "Feedback Saved SuccessFully"
    draft = fake_db.calls_to("usp_InsertSupFeedbacks")[0]
    assert (draft["Status"], draft["SoftSkills"], draft["SupervisorId"]) == ("0", "0", "12345")

    incomplete = user_client.post("/api/feedback/supervisor", json={**body, "submit": True})
    assert incomplete.json()["detail"] == "Please provide rating for Soft skills of Trainee One"

    stranger = {**body, "rows": [{**ROW, "subordinateId": "999"}]}
    assert user_client.post("/api/feedback/supervisor", json=stranger).status_code == 403

    # No mail template is configured in the fake, so the save succeeds with a warning.
    done = user_client.post("/api/feedback/supervisor", json={**body, "submit": True, "rows": [ROW]}).json()
    assert done["message"] == "Feedback Submitted SuccessFully"
    assert done["warning"]
    assert fake_db.calls_to("usp_InsertSupFeedbacks")[-1]["Status"] == "1"


def test_review_approve(user_client, fake_db):
    fake_db.on("usp_SupervisorFeedback_Details", SUBORDINATES, SAVED)
    form = user_client.get("/api/feedback/review", params={"session_id": "S1", "supervisor": "Sup Name-54321"}).json()
    assert form["header"]["employeeId"] == "54321"
    assert form["columns"][:2] == ["Trainee's Employee Id", "Trainee's Name"]
    assert form["rows"][0][:3] == ["777", "Trainee One", "QA"]

    body = {"sessionId": "S1", "courseId": "C1", "supervisor": "Sup Name-54321", "remarks": "", "approve": True}
    assert user_client.post("/api/feedback/review", json=body).json()["detail"] == "Please provide Remarks"
    assert user_client.post("/api/feedback/review", json={**body, "remarks": "ok"}).status_code == 200
    update = fake_db.calls_to("usp_UpdatePMRemarks")[0]
    assert (update["Status"], update["PM_ID"], update["SupervisorId"]) == (3, "12345", "54321")


# --- reports ----------------------------------------------------------------

RATING_COLUMNS = ["SNo", "SECTION", "MAXS", "SCORE"]


def _rating_table(delivery: str) -> db.Table:
    return db.Table(RATING_COLUMNS, [[1, "Admin", 4.0, "4.50"], [2, "Design", 4.0, "4.00"], [3, "Delivery", 4.0, delivery]])


def test_overall_average_matches_old_page(fake_db):
    scores = {"OverAll": "4.00", "T1": "5.00", "T2": "3.00", "T3": "1.00"}
    fake_db.on("usp_rATINGForOverallrpt_V1", handler=lambda p: db.ProcResult([_rating_table(scores[p["TrainerId"]])]))
    # OverAll and every trainer but the last: (4 + 5 + 3) / 3 trainers
    assert reports.overall_trainer_average("S1", ["T1", "T2", "T3"]) == Decimal(4)
    assert [c["TrainerId"] for c in fake_db.calls_to("usp_rATINGForOverallrpt_V1")] == ["OverAll", "T1", "T2"]


def test_overall_report_and_pdf(admin_client, fake_db):
    fake_db.on("usp_getFeedbackForOverallrpt_V1", db.Table(["S_No", "QUESTION", "A", "B", "C", "D", "PATTERN"], [[1, "Q", 1, 2, 3, 4, "A"]]))
    fake_db.on("usp_rATINGForOverallrpt_V1", _rating_table("3.25"))
    fake_db.on("usp_getTrainersForSessionid", db.Table(["TrainerID"], [["Anil(111)"]]))
    fake_db.on("usp_getSessionSessionIdWise", db.Table(
        ["PreparedDate", "sessionid", "startdate", "enddate", "max_participants", "Trainer Name"],
        [["01 Feb 2024", "S1", "04 Mar 2024", "05 Mar 2024", 12, "Anil(111)"]]))
    fake_db.on("usp_getFeedbackCommentsForSessionid", db.Table(
        ["Employee_id", "Comments", "CommentsValid"], [["1", "Shown & kept", "1"], ["2", "Hidden", "0"]]))

    params = {"session_id": "S1", "trainer_id": "OverAll"}
    report = admin_client.get("/api/reports/overall", params=params).json()
    assert report["ratingError"] is None
    assert report["rating"]["rows"][2][3] == "3.25"

    pdf = admin_client.get("/api/reports/overall/pdf", params={**params, "course_name": "Java"})
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF")


def test_overall_report_without_answers(admin_client, fake_db):
    response = admin_client.get("/api/reports/overall", params={"session_id": "S1", "trainer_id": "T1"})
    assert (response.status_code, response.json()["detail"]) == (404, reports.NO_FEEDBACK)


def test_defaulters_hide_completed(admin_client, fake_db):
    fake_db.on("usp_getFeedbackBalanceEmpIdForSessionid", db.Table(
        ["Employee ID", "Employee Name", "Pending FeedBack"], [["1", "A", 0], ["2", "B", 2], ["3", "C", -1]]))
    rows = admin_client.get("/api/reports/defaulters", params={"session_id": "S1"}).json()["rows"]
    assert rows == [["2", "B", 2]]


ANSWERS = db.Table(
    ["Question No", "QUESTION", "Response", "Session", "Employee Id", "Employee Name"],
    [
        ["21", "Administration", 5, "S1", "1", "Ann"],
        ["23", "Delivery", 4, "S1", "1", "Ann"],
        ["21", "Administration", 3, "S1", "2", "Bob"],
        ["23", "Delivery", None, "S1", "2", "Bob"],
    ],
)
COMMENTS = db.Table(["Employee_id", "Comments", "CommentsValid", "TrainerId"], [["1", "nice", "1", "T1"], ["2", "other", "1", "T2"]])


def test_individual_report_all_employees(admin_client, fake_db):
    fake_db.on("usp_getFeedbackdetQuesansEmployeeid_V1", ANSWERS)
    fake_db.on("usp_getFeedbackCommentsForSessionid", COMMENTS)
    fake_db.on("usp_getFeedbackBalanceEmpIdForSessionid", db.Table(
        ["Employee ID", "Employee Name", "Pending FeedBack"], [["2", "Bob", 1], ["3", "Cy", 2]]))
    params = {"session_id": "S1", "trainer_id": "T1", "employee_id": "All"}
    report = admin_client.get("/api/reports/individual", params=params).json()
    assert report["columns"] == ["SessionId", "Trainer", "Employee Id", "Employee Name", "Q21", "Q23", "Comments"]
    assert report["rows"] == [
        ["S1", "T1", "1", "Ann", "5", "4", "nice"],
        ["S1", "T1", "2", "Bob", "3", "", None],
        ["S1", "T1", "3", "Cy", "0", "0", ""],
    ]
    assert [q["number"] for q in report["questions"]] == ["21", "23"]
    assert report["questions"][0]["description"].startswith("Please rate the training execution")

    excel = admin_client.get("/api/reports/individual/excel", params={**params, "course_name": "Java"})
    sheet = load_workbook(BytesIO(excel.content)).active
    assert [c.value for c in sheet[1]][:2] == ["Session", "Java"]
    assert [c.value for c in sheet[5]][:4] == ["SessionId", "Trainer", "Employee Id", "Employee Name"]


def test_individual_report_single_employee(admin_client, fake_db):
    fake_db.on("usp_getFeedbackdetQuesansEmployeeid_V1", db.Table(ANSWERS.columns, ANSWERS.rows[2:]))
    fake_db.on("usp_getFeedbackCommentsForSessionid", COMMENTS)
    params = {"session_id": "S1", "trainer_id": "T2", "employee_id": "2"}
    report = admin_client.get("/api/reports/individual", params=params).json()
    assert report["columns"][-1] == "Comments"
    assert report["rows"] == [["21", "Administration", 3, "S1", "2", "Bob", "other"]]


def test_effectiveness_excel_drops_serial_number(admin_client, fake_db):
    fake_db.on("usp_getTrnEffectivenessReport", db.Table(["serial number", "course"], [[1, "Java"]]))
    excel = admin_client.get("/api/reports/effectiveness/excel", params={"session_id": "S1"})
    assert [[c.value for c in row] for row in load_workbook(BytesIO(excel.content)).active] == [["course"], ["Java"]]


def test_comment_approval(admin_client, fake_db):
    body = {"sessionId": "S1", "comments": [{"employeeId": "1", "display": True}, {"employeeId": "2", "display": False}]}
    assert admin_client.put("/api/reports/comments", json=body).status_code == 200
    assert [c["Commentsvalid"] for c in fake_db.calls_to("usp_UpdateCommentsvaildDetails")] == [1, 0]


# --- Excel data upload ------------------------------------------------------


def _workbook() -> bytes:
    buffer = BytesIO()
    Workbook().save(buffer)
    return buffer.getvalue()


def test_upload_saves_file_calls_import_and_cleans_up(admin_client, fake_db, tmp_path, monkeypatch):
    monkeypatch.setattr(get_settings(), "upload_dir", str(tmp_path))
    seen = {}

    def importer(params):
        seen.update(params)
        seen["existed"] = (tmp_path / params["FILEPATH"].split("\\")[-1].split("/")[-1]).exists()
        return db.ProcResult([db.Table(["n"], [[14]])])

    fake_db.on("InsertCertificationsData", handler=importer)
    form = {"upload_type": "Certifications", "month": "3", "year": "2024"}
    result = admin_client.post("/api/skillport/upload", data=form, files={"file": ("c.xlsx", _workbook())}).json()
    assert result["message"] == "Totally 14 Record(s) saved successfully"
    assert (seen["EMPID"], seen["MONTH"], seen["YEAR"], seen["existed"]) == ("12345", "3", "2024", True)
    assert seen["FILEPATH"].endswith(".xls")
    assert list(tmp_path.iterdir()) == []


def test_upload_refuses_when_month_already_loaded(admin_client, fake_db, tmp_path, monkeypatch):
    monkeypatch.setattr(get_settings(), "upload_dir", str(tmp_path))
    fake_db.on("CheckRecordDuplicate", db.Table(["EmployeeID"], [["1"]]))
    form = {"upload_type": "Learning Plan for PMS", "month": "3", "year": "2024"}
    result = admin_client.post("/api/skillport/upload", data=form, files={"file": ("c.xls", b"x")}).json()
    assert result["saved"] == 0 and result["existing"]["rows"] == [["1"]]
    assert fake_db.calls_to("InsertSkillPortData") == []


def test_upload_validation_and_delete(admin_client, fake_db):
    form = {"upload_type": "Certifications", "month": "3", "year": "2024"}
    bad = admin_client.post("/api/skillport/upload", data=form, files={"file": ("c.txt", b"x")})
    assert bad.json()["detail"] == "The file should be in excel format only."
    future = {**form, "month": "12", "year": str(date.today().year)}
    if date.today().month < 12:
        late = admin_client.post("/api/skillport/upload", data=future, files={"file": ("c.xls", b"x")})
        assert late.json()["detail"] == "Please check the month of your request."

    fake_db.responses["DeleteCertificationsDuplicateRecords"] = 9
    body = {"uploadType": "Certifications", "month": 3, "year": 2024}
    assert admin_client.post("/api/skillport/delete", json=body).json()["deleted"] == 9
